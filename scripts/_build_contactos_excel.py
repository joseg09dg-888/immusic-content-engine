"""Convierte docs/libro/contactos_artistas_afiliados.csv a un .xlsx con
color por nivel de confianza, para el outreach de afiliados."""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
CSV_PATH = ROOT / "docs" / "libro" / "contactos_artistas_afiliados.csv"
OUT_PATH = ROOT / "docs" / "libro" / "Contactos_Artistas_Afiliados_IM_Music.xlsx"

VIOLETA = "5E17EB"
BLANCO = "FFFFFF"
ALTA = "C6EFCE"
MEDIA = "FFEB9C"
BAJA = "FFC7CE"

CONF_FILL = {"Alta": ALTA, "Media": MEDIA, "Baja": BAJA}


def main():
    with CSV_PATH.open(encoding="utf-8") as f:
        rows = list(csv.reader(f))

    header, data = rows[0], rows[1:]

    wb = Workbook()
    ws = wb.active
    ws.title = "Contactos"

    thin = Side(style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for j, h in enumerate(header, start=1):
        cell = ws.cell(row=1, column=j, value=h)
        cell.font = Font(bold=True, color=BLANCO, name="Calibri", size=11)
        cell.fill = PatternFill("solid", fgColor=VIOLETA)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 22

    conf_idx = header.index("Confianza") + 1

    for i, row in enumerate(data, start=2):
        for j, val in enumerate(row, start=1):
            cell = ws.cell(row=i, column=j, value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            cell.font = Font(name="Calibri", size=10)
        conf_val = row[conf_idx - 1].strip()
        fill_color = CONF_FILL.get(conf_val)
        if fill_color:
            ws.cell(row=i, column=conf_idx).fill = PatternFill("solid", fgColor=fill_color)
        ws.row_dimensions[i].height = 90

    widths = {
        "A": 22,  # Artista
        "B": 26,  # Agencia/Manager
        "C": 30,  # Correo/Contacto
        "D": 18,  # Tipo
        "E": 40,  # URL fuente
        "F": 12,  # Confianza
        "G": 60,  # Notas
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    ws.freeze_panes = "A2"

    # Leyenda
    legend_row = len(data) + 3
    ws.cell(row=legend_row, column=1, value="Leyenda de confianza:").font = Font(bold=True, size=10)
    labels = [("Alta", "Confirmado por fuente oficial directa"),
              ("Media", "Fuente de terceros confiable (agencia real), no exclusiva del artista"),
              ("Baja", "NO verificado directamente — requiere confirmación manual antes de usar")]
    for k, (label, desc) in enumerate(labels):
        r = legend_row + 1 + k
        c = ws.cell(row=r, column=1, value=label)
        c.fill = PatternFill("solid", fgColor=CONF_FILL[label])
        c.font = Font(bold=True, size=9)
        ws.cell(row=r, column=2, value=desc).font = Font(size=9, italic=True)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(OUT_PATH))
    print(f"OK: {OUT_PATH} ({OUT_PATH.stat().st_size // 1024}KB)")


if __name__ == "__main__":
    main()
